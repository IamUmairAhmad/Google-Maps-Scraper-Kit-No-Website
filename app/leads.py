"""Classify Maps listings (no website / has website / closed), filter them, and write CSV or Excel."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Iterable, Mapping
from urllib.parse import parse_qs, urlparse

LEAD_COLUMNS = (
    "name",
    "phone",
    "email",
    "website",
    "category",
    "address",
    "rating",
    "review_count",
    "status",
    "result",
    "maps_url",
    "niche",
    "location",
)

RESULT_NO_WEBSITE = "No website"
RESULT_HAS_WEBSITE = "Has website"
RESULT_CLOSED = "Closed"

# The four options on the page. "checked" is every listing that was opened.
FILTER_KEYS = ("checked", "no_website", "has_website", "closed")
_RESULT_FOR_FILTER = {
    "no_website": RESULT_NO_WEBSITE,
    "has_website": RESULT_HAS_WEBSITE,
    "closed": RESULT_CLOSED,
}
_FILENAME_FOR_FILTER = {
    "checked": "all-checked",
    "no_website": "no-website",
    "has_website": "has-website",
    "closed": "closed",
}

_EMPTY_WEBSITE = {"", "n/a", "na", "-", "none", "null", "undefined"}
_FORMULA_PREFIX = ("=", "+", "-", "@")
_PROFILE_HOSTS = (
    "facebook.com",
    "fb.com",
    "fb.me",
    "instagram.com",
    "twitter.com",
    "x.com",
    "tiktok.com",
    "linkedin.com",
    "youtube.com",
    "youtu.be",
    "pinterest.com",
    "linktr.ee",
    "linktree.com",
)


def classify_status(panel_text: str) -> str:
    """Map the place header to active, temporarily closed, or permanently closed.

    "Closed · Opens 9 AM" is still an operating business. Only the closed-down
    labels drop a listing.
    """
    text = (panel_text or "").lower()
    if "permanently closed" in text:
        return "permanently closed"
    if "temporarily closed" in text:
        return "temporarily closed"
    return "active"


def _hostname(url: str) -> str:
    candidate = url if "://" in url else f"https://{url}"
    host = (urlparse(candidate).hostname or "").lower()
    if host.startswith("www."):
        host = host[4:]
    return host


def unwrap_google_redirect(url: str) -> str:
    text = (url or "").strip()
    parsed = urlparse(text)
    host = _hostname(text)
    if host in {"google.com", "maps.google.com"} and parsed.path.startswith("/url"):
        target = parse_qs(parsed.query).get("q", [""])[0].strip()
        return target or text
    return text


def _is_profile_host(host: str) -> bool:
    return any(host == name or host.endswith("." + name) for name in _PROFILE_HOSTS)


def has_website(value: str | None) -> bool:
    """True when Maps lists a real site. Social profiles do not count."""
    text = unwrap_google_redirect(value).strip()
    if text.lower() in _EMPTY_WEBSITE:
        return False
    return not _is_profile_host(_hostname(text))


def business_website(value: str | None) -> str:
    text = unwrap_google_redirect(value).strip()
    if not has_website(text):
        return ""
    return text


def review_count_from_panel(panel_text: str) -> str:
    header = (panel_text or "")[:900]
    # "(530)" is the review count. "(207) 801-9191" is a phone, so skip it.
    match = re.search(r"\((\d[\d,]*)\)(?!\s*\d)", header) or re.search(
        r"(\d[\d,]*)\s+reviews\b", header, re.IGNORECASE
    )
    if not match:
        return ""
    return match.group(1).replace(",", "")


def is_active(status: str | None) -> bool:
    return (status or "").strip().lower() == "active"


def classify_result(row: Mapping[str, str]) -> str:
    """One label per listing: Closed, No website, or Has website."""
    if not is_active(row.get("status")):
        return RESULT_CLOSED
    if has_website(row.get("website")):
        return RESULT_HAS_WEBSITE
    return RESULT_NO_WEBSITE


def keep_lead(row: Mapping[str, str]) -> bool:
    """True when the business is operating and Maps has no website for it."""
    return classify_result(row) == RESULT_NO_WEBSITE


def filter_rows(rows: Iterable[Mapping[str, object]], key: str) -> list:
    """Rows for one of the four options. "checked" returns every row."""
    if key not in FILTER_KEYS:
        raise ValueError(f"Unknown filter: {key}")
    rows = list(rows)
    if key == "checked":
        return rows
    label = _RESULT_FOR_FILTER[key]
    return [row for row in rows if classify_result(row) == label]  # type: ignore[arg-type]


def count_results(rows: Iterable[Mapping[str, object]]) -> dict:
    counts = {"checked": 0, "found": 0, "skipped_website": 0, "skipped_closed": 0}
    for row in rows:
        counts["checked"] += 1
        label = classify_result(row)  # type: ignore[arg-type]
        if label == RESULT_NO_WEBSITE:
            counts["found"] += 1
        elif label == RESULT_HAS_WEBSITE:
            counts["skipped_website"] += 1
        else:
            counts["skipped_closed"] += 1
    return counts


def with_result(row: Mapping[str, object]) -> dict:
    """Copy of the row with its `result` column filled in."""
    out = dict(row)
    out["result"] = classify_result(row)  # type: ignore[arg-type]
    return out


def csv_cell(value: object) -> str:
    text = "" if value is None else str(value)
    if text.startswith(_FORMULA_PREFIX):
        return "'" + text
    return text


def rows_to_csv(rows: Iterable[Mapping[str, object]]) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=LEAD_COLUMNS, lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({column: csv_cell(row.get(column, "")) for column in LEAD_COLUMNS})
    return buffer.getvalue()


def clickable_maps_url(url: str | None) -> str:
    """A Maps link safe to put behind a name. Empty when there is no usable link."""
    text = (url or "").strip()
    if not text.lower().startswith(("https://", "http://")):
        return ""
    if len(text) <= 2000:
        return text
    # Very long place URLs break some spreadsheet apps. The CID form is short.
    match = re.search(r"!1s0x[0-9a-f]+:0x([0-9a-f]+)", text)
    if match:
        return f"https://www.google.com/maps?cid={int(match.group(1), 16)}"
    return ""


def rows_to_xlsx(rows: Iterable[Mapping[str, object]]) -> bytes:
    """Excel file where each Name is a clickable link to its Google Maps listing."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(list(LEAD_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    link_font = Font(color="0563C1", underline="single")
    widths = [len(column) for column in LEAD_COLUMNS]
    for row_number, row in enumerate(rows, start=2):
        for column_number, column in enumerate(LEAD_COLUMNS, start=1):
            value = row.get(column, "")
            text = "" if value is None else str(value)
            cell = sheet.cell(row=row_number, column=column_number, value=text)
            cell.data_type = "s"  # always text, never a formula
            widths[column_number - 1] = max(widths[column_number - 1], min(len(text), 60))
        link = clickable_maps_url(str(row.get("maps_url", "") or ""))
        name_cell = sheet.cell(row=row_number, column=1)
        if link and name_cell.value:
            name_cell.hyperlink = link
            name_cell.font = link_font
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width + 2
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def slug(value: str, fallback: str = "search") -> str:
    text = re.sub(r"[^a-z0-9]+", "-", (value or "").strip().lower()).strip("-")
    return text or fallback


def result_filename(key: str, niche: str, location: str, extension: str = "csv") -> str:
    if key not in FILTER_KEYS:
        raise ValueError(f"Unknown filter: {key}")
    return f"{_FILENAME_FOR_FILTER[key]}-{slug(niche)}-{slug(location)}.{extension}"


def write_results(rows: list, niche: str, location: str, folder: str | Path = "results") -> list:
    """Save one CSV per option (the all-checked file always, the others when not empty)."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    saved = []
    for key in FILTER_KEYS:
        subset = filter_rows(rows, key)
        if not subset and key != "checked":
            continue
        path = folder / result_filename(key, niche, location)
        path.write_text(rows_to_csv(subset), encoding="utf-8")
        saved.append(path)
    return saved
